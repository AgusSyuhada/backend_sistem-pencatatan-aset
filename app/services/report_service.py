import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class ReportService:
    @staticmethod
    def generate_cycle_excel(data: dict) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "Laporan Siklus"

        std_font = Font(name="Arial", size=10)
        bold_font = Font(name="Arial", size=10, bold=True)

        thin_border_side = Side(style="thin")
        medium_border_side = Side(style="medium")

        all_border = Border(
            left=thin_border_side,
            right=thin_border_side,
            top=thin_border_side,
            bottom=thin_border_side,
        )

        header_bottom_border = Border(bottom=medium_border_side)

        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
        top_align_wrap = Alignment(vertical="top", wrap_text=True)

        ws.merge_cells("A1:T1")
        cell_a1 = ws["A1"]
        cell_a1.value = "LAPORAN SIKLUS"
        cell_a1.font = bold_font
        cell_a1.alignment = center_align

        cycle_num = data.get('cycle', 0)
        cycle_label = str(cycle_num)
        
        if cycle_num == 1:
            cycle_label = "JANUARI - APRIL"
        elif cycle_num == 2:
            cycle_label = "MEI - AGUSTUS"
        elif cycle_num == 3:
            cycle_label = "SEPTEMBER - DESEMBER"

        ws.merge_cells("A2:T2")
        cell_a2 = ws["A2"]
        cell_a2.value = f"PERIODE {cycle_label} TAHUN {data.get('year', '')}"
        cell_a2.font = bold_font
        cell_a2.alignment = center_align

        for col in range(1, 21):
            cell = ws.cell(row=2, column=col)
            cell.border = header_bottom_border

        stats_main = ["Statistik Aset"] + [
            f"Total Aset : {data.get('total_assets', 0)}",
            f"Selesai : {data.get('cycled_assets', 0)}",
            f"Tersisa : {data.get('pending_assets', 0)}",
            f"Nilai Aset : {data.get('total_asset_value', 0)}",
        ]

        cond_raw = data.get("dist_condition", [])
        stats_cond = ["Distribusi Kondisi"] + [
            f"{item['name']} : {item['count']} Buah" for item in cond_raw
        ]

        area_raw = data.get("dist_area", [])
        stats_area = ["Distribusi Area"] + [
            f"{item['name']} : {item['count']} Lokasi" for item in area_raw
        ]

        max_rows = max(len(stats_main), len(stats_cond), len(stats_area))
        start_stat_row = 4

        for i in range(max_rows):
            current_row = start_stat_row + i

            if i < len(stats_main):
                ws.merge_cells(
                    start_row=current_row,
                    start_column=1,
                    end_row=current_row,
                    end_column=6,
                )
                cell = ws.cell(row=current_row, column=1)
                cell.value = stats_main[i]

                cell.font = bold_font if i == 0 else std_font
                cell.alignment = left_align

            if i < len(stats_cond):
                ws.merge_cells(
                    start_row=current_row,
                    start_column=7,
                    end_row=current_row,
                    end_column=12,
                )
                cell = ws.cell(row=current_row, column=7)
                cell.value = stats_cond[i]
                cell.font = bold_font if i == 0 else std_font
                cell.alignment = left_align

            if i < len(stats_area):
                ws.merge_cells(
                    start_row=current_row,
                    start_column=13,
                    end_row=current_row,
                    end_column=20,
                )
                cell = ws.cell(row=current_row, column=13)
                cell.value = stats_area[i]
                cell.font = bold_font if i == 0 else std_font
                cell.alignment = left_align

        table_header_row = start_stat_row + max_rows + 1

        headers = [
            "NO",
            "NO ASSET",
            "HBM",
            "NAMA ASET",
            "NAMA TEAM",
            "NILAI ASET (Rp)",
            "COST CENTER",
            "NOMOR SERI",
            "MODEL/TYPE",
            "MANUFACTURE",
            "GPS KOORDINAT",
            "KONDISI",
            "KODE LOKASI SAP",
            "AREA",
            "LOCATION",
            "LOKASI SPESIFIK",
            "KETERANGAN",
            "HASIL INVENTORY",
            "TANGGAL INVENTORY",
            "PIC TEAM FAV",
        ]

        for col_idx, header_text in enumerate(headers, 1):
            cell = ws.cell(row=table_header_row, column=col_idx, value=header_text)
            cell.font = bold_font
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.border = all_border

        assets = data.get("assets", [])
        for r_idx, asset in enumerate(assets, 1):
            current_row = table_header_row + r_idx

            row_data = [
                asset.get("index"),
                asset.get("asset_number"),
                asset.get("hbm"),
                asset.get("asset_name"),
                asset.get("team"),
                asset.get("val"),
                asset.get("cc"),
                asset.get("serial"),
                asset.get("model"),
                asset.get("brand"),
                asset.get("gps"),
                asset.get("cond"),
                asset.get("sap"),
                asset.get("area"),
                asset.get("loc"),
                asset.get("spec"),
                asset.get("desc"),
                asset.get("res"),
                asset.get("date"),
                asset.get("pic"),
            ]

            for c_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=current_row, column=c_idx, value=val)
                cell.font = std_font
                cell.border = all_border
                cell.alignment = top_align_wrap

        column_widths = [
            5,
            12,
            12,
            30,
            15,
            15,
            12,
            15,
            15,
            15,
            18,
            12,
            12,
            15,
            15,
            25,
            25,
            12,
            12,
            20,
        ]

        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.read()
